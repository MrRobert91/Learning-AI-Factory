import io
import zipfile

import pytest
from factory_agents.agents.ideation import AgentEvent
from factory_api.ideation_sources import (
    MAX_SOURCE_BYTES,
    SourceIngestionError,
    capture_bytes,
    capture_url,
)
from pptx import Presentation
from reportlab.pdfgen import canvas


def _brief_event():
    return AgentEvent(
        kind="brief",
        content="Brief propuesto",
        payload={
            "working_title": "Curso basado en fuentes",
            "topic": "RAG",
            "audience": "Desarrolladores",
            "level": "intermedio",
            "duration_spec": {
                "preset": "microvideo",
                "module_count": 1,
                "videos_per_module": 1,
                "target_minutes_per_video": 1,
            },
        },
    )


def _docx_bytes() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "word/document.xml",
            (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
                '2006/main"><w:body><w:p><w:r><w:t>Texto DOCX trazable</w:t>'
                "</w:r></w:p></w:body></w:document>"
            ),
        )
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Texto PPTX trazable"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output)
    document.drawString(72, 720, "Texto PDF trazable")
    document.save()
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content", "kind", "needle"),
    [
        ("fuente.md", b"# Texto Markdown trazable", "markdown", "Markdown"),
        ("fuente.docx", _docx_bytes(), "docx", "DOCX"),
        ("fuente.pptx", _pptx_bytes(), "pptx", "PPTX"),
        ("fuente.pdf", _pdf_bytes(), "pdf", "PDF"),
    ],
    ids=["markdown", "docx", "pptx", "pdf"],
)
def test_supported_documents_are_detected_and_extracted(filename, content, kind, needle):
    captured = capture_bytes(content, filename=filename)
    assert captured.kind == kind
    assert needle in captured.extracted_text
    assert len(captured.sha256) == 64


def test_url_capture_blocks_private_addresses():
    with pytest.raises(SourceIngestionError, match="privada|reservada"):
        capture_url("http://127.0.0.1/private")


def test_invalid_size_and_dangerous_extension_mismatch_are_rejected():
    with pytest.raises(SourceIngestionError, match="25 MB"):
        capture_bytes(b"x" * (MAX_SOURCE_BYTES + 1), filename="too-large.txt")
    with pytest.raises(SourceIngestionError, match="extensión"):
        capture_bytes(_pdf_bytes(), filename="disguised.txt")


def test_source_is_frozen_into_brief_and_project(auth_client, monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test")
    seen = {}

    def fake_turn(client, model, history, **kwargs):
        seen.update(kwargs)
        return [_brief_event()]

    monkeypatch.setattr("factory_api.routers.ideation.run_ideation_turn", fake_turn)
    monkeypatch.setattr(
        "factory_api.routers.ideation.get_llm_client", lambda key: object()
    )

    draft = auth_client.post(
        "/api/ideation/draft",
        json={"idea": "curso de RAG", "research_mode": "provided_only"},
    )
    assert draft.status_code == 201
    session_id = draft.json()["id"]

    uploaded = auth_client.post(
        f"/api/ideation/{session_id}/sources/file",
        files={"file": ("notes.txt", b"RAG usa recuperacion y generacion.", "text/plain")},
    )
    assert uploaded.status_code == 201
    source = uploaded.json()
    assert source["status"] == "ready"
    assert source["kind"] == "text"
    assert auth_client.get(f"/api/ideation/sources/{source['id']}/text").status_code == 200

    started = auth_client.post(f"/api/ideation/{session_id}/start/stream")
    assert started.status_code == 200
    session = auth_client.get(f"/api/ideation/{session_id}").json()
    assert session["brief"]["research_mode"] == "provided_only"
    assert session["brief"]["source_ids"] == [source["id"]]
    assert seen["research_mode"] == "provided_only"
    assert seen["source_documents"][0]["text"].startswith("RAG")

    finalized = auth_client.post(f"/api/ideation/{session_id}/finalize")
    assert finalized.status_code == 200
    project = finalized.json()
    assert project["research_mode"] == "provided_only"
    assert [item["id"] for item in project["sources"]] == [source["id"]]


def test_failed_source_stays_visible_without_breaking_session(auth_client):
    draft = auth_client.post(
        "/api/ideation/draft",
        json={"idea": "curso", "research_mode": "provided_plus_web"},
    ).json()
    uploaded = auth_client.post(
        f"/api/ideation/{draft['id']}/sources/file",
        files={"file": ("malware.pdf", b"not a pdf\x00", "application/pdf")},
    )
    assert uploaded.status_code == 201
    assert uploaded.json()["status"] == "failed"
    session = auth_client.get(f"/api/ideation/{draft['id']}").json()
    assert session["sources"][0]["status"] == "failed"
