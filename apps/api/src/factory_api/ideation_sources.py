"""Safe, reproducible ingestion for ideation source files and URLs."""

import hashlib
import ipaddress
import json
import re
import socket
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from pptx import Presentation
from pypdf import PdfReader

MAX_SOURCE_BYTES = 25 * 1024 * 1024
MAX_SOURCES = 10
MAX_REDIRECTS = 5
SOURCE_TIMEOUT_SECONDS = 25


class SourceIngestionError(ValueError):
    pass


@dataclass
class CapturedSource:
    kind: str
    media_type: str
    name: str
    content: bytes
    extracted_text: str
    metadata: dict = field(default_factory=dict)
    original_url: str | None = None
    final_url: str | None = None

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.content).hexdigest()


def _normalize_text(value: str) -> str:
    value = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _zip_members(content: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            return set(archive.namelist())
    except zipfile.BadZipFile:
        return set()


def detect_kind(content: bytes, filename: str, content_type: str = "") -> tuple[str, str]:
    suffix = Path(filename).suffix.casefold()
    if content.startswith(b"%PDF-"):
        if suffix and suffix != ".pdf":
            raise SourceIngestionError("La extensión no coincide con el PDF detectado")
        return "pdf", "application/pdf"

    members = _zip_members(content)
    if "word/document.xml" in members:
        if suffix and suffix != ".docx":
            raise SourceIngestionError("La extensión no coincide con el DOCX detectado")
        return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if "ppt/presentation.xml" in members:
        if suffix and suffix != ".pptx":
            raise SourceIngestionError("La extensión no coincide con el PPTX detectado")
        return "pptx", (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

    sample = content[:4096]
    if b"\x00" not in sample:
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = ""
        if decoded:
            lowered_type = content_type.casefold()
            if "html" in lowered_type or re.search(
                r"<(?:!doctype\s+html|html|head|body)(?:\s|>)",
                decoded[:2000],
                re.IGNORECASE,
            ):
                return "html", "text/html"
            if suffix in {".md", ".markdown"}:
                return "markdown", "text/markdown"
            if suffix == ".txt" or not suffix:
                return "text", "text/plain"
    raise SourceIngestionError(
        "Formato no admitido o contenido incompatible; usa PDF, DOCX, PPTX, Markdown o TXT"
    )


def _extract_pdf(content: bytes) -> tuple[str, dict]:
    try:
        reader = PdfReader(BytesIO(content))
        pages = [
            f"[Página {index}]\n{_normalize_text(page.extract_text() or '')}"
            for index, page in enumerate(reader.pages, 1)
        ]
    except Exception as exc:
        raise SourceIngestionError(f"No se pudo leer el PDF: {exc}") from exc
    text = "\n\n".join(page for page in pages if page.split("\n", 1)[-1].strip())
    if not text:
        raise SourceIngestionError(
            "El PDF no contiene una capa de texto; los documentos escaneados requieren OCR"
        )
    return text, {"pages": len(reader.pages)}


def _extract_docx(content: bytes) -> tuple[str, dict]:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            xml = archive.read("word/document.xml")
    except Exception as exc:
        raise SourceIngestionError(f"No se pudo leer el DOCX: {exc}") from exc
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        raise SourceIngestionError(f"El DOCX contiene XML inválido: {exc}") from exc
    paragraphs = []
    for paragraph in (node for node in root.iter() if node.tag.endswith("}p")):
        text = "".join(
            node.text or "" for node in paragraph.iter() if node.tag.endswith("}t")
        )
        paragraphs.append(_normalize_text(text))
    text = "\n".join(item for item in paragraphs if item)
    if not text:
        raise SourceIngestionError("El DOCX no contiene texto extraíble")
    return f"[Sección documento]\n{text}", {"paragraphs": len(paragraphs)}


def _extract_pptx(content: bytes) -> tuple[str, dict]:
    try:
        presentation = Presentation(BytesIO(content))
    except Exception as exc:
        raise SourceIngestionError(f"No se pudo leer el PPTX: {exc}") from exc
    rendered: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        texts = [
            _normalize_text(shape.text)
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            rendered.append(f"[Diapositiva {index}]\n" + "\n".join(texts))
    if not rendered:
        raise SourceIngestionError("El PPTX no contiene texto extraíble")
    return "\n\n".join(rendered), {"slides": len(presentation.slides)}


def _extract_html(content: bytes) -> tuple[str, dict]:
    try:
        html = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        html = content.decode("latin-1")
    soup = BeautifulSoup(html, "html.parser")
    title = _normalize_text(soup.title.get_text(" ")) if soup.title else ""
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside"]):
        tag.decompose()
    text = _normalize_text(soup.get_text("\n"))
    if not text:
        raise SourceIngestionError("La página HTML no contiene texto extraíble")
    return f"[Sección página]\n{text}", {"title": title}


def capture_bytes(
    content: bytes,
    *,
    filename: str,
    content_type: str = "",
    original_url: str | None = None,
    final_url: str | None = None,
) -> CapturedSource:
    if not content:
        raise SourceIngestionError("La fuente está vacía")
    if len(content) > MAX_SOURCE_BYTES:
        raise SourceIngestionError("La fuente supera el límite de 25 MB")
    kind, media_type = detect_kind(content, filename, content_type)
    if kind == "pdf":
        text, metadata = _extract_pdf(content)
    elif kind == "docx":
        text, metadata = _extract_docx(content)
    elif kind == "pptx":
        text, metadata = _extract_pptx(content)
    elif kind == "html":
        text, metadata = _extract_html(content)
    else:
        try:
            text = _normalize_text(content.decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise SourceIngestionError("El texto debe estar codificado en UTF-8") from exc
        metadata = {"lines": len(text.splitlines())}
        if not text:
            raise SourceIngestionError("La fuente no contiene texto")
    safe_name = Path(filename).name[:255] or f"fuente.{kind}"
    return CapturedSource(
        kind=kind,
        media_type=media_type,
        name=safe_name,
        content=content,
        extracted_text=text,
        metadata=metadata,
        original_url=original_url,
        final_url=final_url,
    )


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise SourceIngestionError("La URL debe usar HTTP o HTTPS")
    hostname = parsed.hostname.rstrip(".").casefold()
    if hostname == "localhost" or hostname.endswith(".localhost"):
        raise SourceIngestionError("La URL apunta a una red local bloqueada")
    try:
        addresses = {
            item[4][0]
            for item in socket.getaddrinfo(
                hostname,
                parsed.port or (443 if parsed.scheme == "https" else 80),
                type=socket.SOCK_STREAM,
            )
        }
    except socket.gaierror as exc:
        raise SourceIngestionError("No se pudo resolver el dominio de la URL") from exc
    if not addresses:
        raise SourceIngestionError("No se pudo resolver el dominio de la URL")
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if not ip.is_global:
            raise SourceIngestionError("La URL apunta a una IP privada o reservada")


def capture_url(url: str) -> CapturedSource:
    current = url.strip()
    original = current
    headers = {"User-Agent": "AI Learning Factory source capture/1.0"}
    with httpx.Client(timeout=SOURCE_TIMEOUT_SECONDS, follow_redirects=False) as client:
        for redirect_count in range(MAX_REDIRECTS + 1):
            _validate_public_url(current)
            try:
                with client.stream("GET", current, headers=headers) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise SourceIngestionError("La redirección no incluye destino")
                        if redirect_count >= MAX_REDIRECTS:
                            raise SourceIngestionError("La URL supera el límite de redirecciones")
                        current = urljoin(current, location)
                        continue
                    response.raise_for_status()
                    chunks: list[bytes] = []
                    size = 0
                    for chunk in response.iter_bytes():
                        size += len(chunk)
                        if size > MAX_SOURCE_BYTES:
                            raise SourceIngestionError("La fuente supera el límite de 25 MB")
                        chunks.append(chunk)
                    content = b"".join(chunks)
                    content_type = response.headers.get("content-type", "")
            except SourceIngestionError:
                raise
            except (httpx.HTTPError, OSError) as exc:
                raise SourceIngestionError(f"No se pudo capturar la URL: {exc}") from exc
            break
        else:
            raise SourceIngestionError("La URL supera el límite de redirecciones")

    parsed = urlparse(current)
    filename = Path(unquote(parsed.path)).name or "pagina.html"
    if "pdf" in content_type.casefold() and not filename.casefold().endswith(".pdf"):
        filename += ".pdf"
    captured = capture_bytes(
        content,
        filename=filename,
        content_type=content_type,
        original_url=original,
        final_url=current,
    )
    if captured.kind not in {"html", "pdf"}:
        raise SourceIngestionError("Las URL públicas solo pueden apuntar a HTML o PDF")
    if captured.kind == "html" and captured.metadata.get("title"):
        captured.name = str(captured.metadata["title"])[:255]
    return captured


def store_captured_source(
    captured: CapturedSource,
    *,
    data_dir: Path,
    session_id: str,
    source_id: str,
) -> str:
    suffixes = {
        "pdf": ".pdf",
        "docx": ".docx",
        "pptx": ".pptx",
        "markdown": ".md",
        "text": ".txt",
        "html": ".html",
    }
    relative = Path("sources") / session_id / source_id / (
        "original" + suffixes[captured.kind]
    )
    target = data_dir / relative
    target.parent.mkdir(parents=True, exist_ok=False)
    target.write_bytes(captured.content)
    (target.parent / "extracted.txt").write_text(captured.extracted_text, encoding="utf-8")
    (target.parent / "metadata.json").write_text(
        json.dumps(captured.metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return relative.as_posix()
