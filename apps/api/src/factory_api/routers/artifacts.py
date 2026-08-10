import hashlib
import io
import json
import logging
import re
import shutil
import uuid
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Annotated

from factory_agents.tools.images import (
    CONSISTENCY_PROMPT,
    DEFAULT_IMAGE_MODEL,
    DEFAULT_IMAGE_STYLE,
    ImageGenerationError,
    consistency_seed,
    generate_image,
    media_extension,
    normalize_automatic_image_layout,
    replace_generated_image,
    resolve_style_prompt,
)
from factory_agents.tools.marp import (
    available_renders,
    inline_local_images,
    legacy_vertical_render_is_current,
    marp_available,
    normalize_marp_canvas,
    render_deck,
    render_manifest_path,
    vertical_theme_path,
)
from factory_agents.tools.palette import (
    DEFAULT_SLIDE_PALETTE,
    apply_slide_palette,
    normalize_palette,
    palette_contrast,
    palette_preset_name,
    palette_warnings,
)
from factory_agents.tools.slide_layout import SlideLayoutError, prepare_slide_layout
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import FileResponse, Response
from pypdf import PdfReader, PdfWriter
from sqlalchemy import select
from sqlalchemy.orm import Session

from factory_api.artifact_versions import (
    add_artifact_version,
    artifact_metadata,
    select_artifact_version,
    select_latest_remaining,
)
from factory_api.auth import CurrentUser
from factory_api.config import get_settings
from factory_api.db import get_db
from factory_api.models import Artifact, Project
from factory_api.pdf_exports import LessonDocument, lessons_pdf
from factory_api.schemas import (
    ArtifactEdit,
    ArtifactRead,
    ArtifactVersionRead,
    SlideImageRegenerate,
    SlidePaletteApply,
    SlidePalettePreview,
)
from factory_api.usage import attach_usage_to_artifact, record_usage

router = APIRouter(prefix="/api", tags=["artifacts"])
logger = logging.getLogger(__name__)

DB = Annotated[Session, Depends(get_db)]

TEXT_FORMATS = {"markdown", "json", "text"}
UPLOADABLE_TYPES = {
    "research_brief": "markdown",
    "course_plan": "json",
    "lesson_content": "markdown",
    "slide_deck": "markdown",
    "teaching_script": "markdown",
}


def _safe_filename(value: str, fallback: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-._")
    return value[:100] or fallback


def _download_response(content: bytes, media_type: str, filename: str) -> Response:
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _selected_artifacts(db: Session, project_id: str, type_: str) -> list[Artifact]:
    return list(
        db.scalars(
            select(Artifact)
            .where(
                Artifact.project_id == project_id,
                Artifact.type == type_,
                Artifact.is_selected.is_(True),
            )
            .order_by(Artifact.title, Artifact.created_at)
        ).all()
    )


def _artifact_documents(artifacts: list[Artifact]) -> list[LessonDocument]:
    settings = get_settings()
    documents = []
    for artifact in artifacts:
        path = settings.data_dir / artifact.path
        if path.is_file():
            documents.append(
                LessonDocument(
                    title=artifact.title or "Leccion",
                    markdown=path.read_text(encoding="utf-8"),
                )
            )
    return documents


def _safe_stored_asset(relative_path: str) -> Path | None:
    if not relative_path:
        return None
    settings = get_settings()
    root = settings.data_dir.resolve()
    candidate = (settings.data_dir / relative_path).resolve()
    if not candidate.is_relative_to(root):
        return None
    return candidate


def _clone_slide_assets(
    markdown: str,
    metadata: dict,
    project_id: str,
    *,
    suffix: str,
) -> tuple[str, dict, Path, str]:
    """Copy a deck's generated images so every artifact version is self-contained."""
    cloned = deepcopy(metadata)
    asset_dir_name = f"slide-assets-{suffix}-{uuid.uuid4().hex[:8]}"
    storage_asset_dir = f"artifacts/{project_id}/{asset_dir_name}"
    output_dir = get_settings().data_dir / storage_asset_dir
    asset_records = list(cloned.get("images", []))
    if isinstance(cloned.get("logo"), dict):
        asset_records.append(cloned["logo"])
    for asset in asset_records:
        stored_path = asset.get("path")
        markdown_path = asset.get("markdown_path")
        if not stored_path or not markdown_path:
            continue
        source = _safe_stored_asset(str(stored_path))
        if source is None or not source.is_file():
            continue
        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / source.name
        shutil.copy2(source, target)
        new_markdown_path = f"{asset_dir_name}/{source.name}"
        markdown = markdown.replace(str(markdown_path), new_markdown_path)
        asset["path"] = f"{storage_asset_dir}/{source.name}"
        asset["markdown_path"] = new_markdown_path
    return markdown, cloned, output_dir, storage_asset_dir


def _remove_slide_assets(metadata: dict, project_id: str) -> None:
    """Best-effort cleanup of the private asset directory owned by one deck version."""
    root = (get_settings().data_dir / "artifacts" / project_id).resolve()
    directories: set[Path] = set()
    asset_records = list(metadata.get("images", []))
    if isinstance(metadata.get("logo"), dict):
        asset_records.append(metadata["logo"])
    for asset in asset_records:
        stored_path = asset.get("path")
        if not stored_path:
            continue
        path = _safe_stored_asset(str(stored_path))
        if path is not None:
            directories.add(path.parent)
    for directory in directories:
        if (
            directory.is_relative_to(root)
            and directory.parent == root
            and directory.name.startswith("slide-assets-")
        ):
            shutil.rmtree(directory, ignore_errors=True)


def _ensure_slide_renders(path: Path) -> dict[str, str]:
    renders = available_renders(path)
    source = path.read_text(encoding="utf-8")
    legacy_vertical = normalize_marp_canvas(source) != source
    if len(renders) < 3 or (
        legacy_vertical and not legacy_vertical_render_is_current(path)
    ):
        render_deck(path)
        renders = available_renders(path)
    return renders


def _prepare_slide_version(
    path: Path,
    markdown: str,
    metadata: dict,
) -> tuple[dict, dict[str, str]]:
    """Validate one immutable deck before its DB row can become selected."""
    prepared_metadata = deepcopy(metadata)
    orientation = str(prepared_metadata.get("orientation", "horizontal"))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(normalize_marp_canvas(markdown), encoding="utf-8")
    try:
        layout = prepare_slide_layout(path, orientation=orientation)
    except SlideLayoutError as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    prepared_metadata["layout_validation"] = layout.metadata
    rendered = render_deck(path) if marp_available() else {}
    if marp_available():
        missing = {"html", "pdf", "pptx"} - set(rendered)
        if missing:
            _cleanup_palette_files([path], [], path.parent.name)
            raise HTTPException(
                status_code=503,
                detail="No se pudieron generar los renders: "
                + ", ".join(sorted(missing)),
            )
    return prepared_metadata, rendered


def _marp_parts(markdown: str) -> tuple[str, str]:
    source = markdown.replace("\r\n", "\n").strip()
    if not source.startswith("---\n"):
        return "", source
    end = source.find("\n---", 4)
    if end == -1:
        return "", source
    after = source.find("\n", end + 4)
    return source[: end + 4], source[after + 1 :] if after != -1 else ""


def _combined_marp(artifacts: list[Artifact]) -> str:
    settings = get_settings()
    frontmatter = ""
    bodies: list[str] = []
    for artifact in artifacts:
        artifact_path = settings.data_dir / artifact.path
        source = artifact_path.read_text(encoding="utf-8")
        source = inline_local_images(source, artifact_path.parent)
        candidate_frontmatter, body = _marp_parts(source)
        if not frontmatter and candidate_frontmatter:
            frontmatter = candidate_frontmatter
        bodies.append(body.strip())
    if not frontmatter:
        frontmatter = "---\nmarp: true\ntheme: default\npaginate: true\n---"
    return frontmatter + "\n\n" + "\n\n---\n\n".join(bodies) + "\n"


def _render_combined_slides(artifacts: list[Artifact], fmt: str) -> bytes:
    settings = get_settings()
    export_dir = settings.data_dir / "tmp" / f"slides-{uuid.uuid4().hex}"
    export_dir.mkdir(parents=True, exist_ok=True)
    source = export_dir / "all-slides.md"
    source.write_text(_combined_marp(artifacts), encoding="utf-8")
    try:
        rendered = render_deck(source)
        target = rendered.get(fmt)
        if target is None:
            raise HTTPException(
                status_code=503,
                detail=f"No se pudo generar el fichero {fmt.upper()}; comprueba Marp CLI",
            )
        return Path(target).read_bytes()
    finally:
        for candidate in export_dir.glob("*"):
            candidate.unlink(missing_ok=True)
        export_dir.rmdir()


def _check_project(db: Session, user_id: str, project_id: str) -> Project:
    project = db.get(Project, project_id)
    if project is None or project.owner_id != user_id:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    return project


def pptx_to_marp(data: bytes) -> str:
    """Convert an uploaded PPTX to Marp Markdown (titles + text, best effort)."""
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    parts = ["---", "marp: true", "theme: default", "paginate: true", "---", ""]
    for i, slide in enumerate(prs.slides):
        if i > 0:
            parts.append("\n---\n")
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except (AttributeError, KeyError):
            pass
        parts.append(f"## {title or f'Slide {i + 1}'}\n")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if title and shape is slide.shapes.title:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    parts.append(f"- {text}")
    return "\n".join(parts) + "\n"


RENDER_MEDIA_TYPES = {
    "html": "text/html",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _check_owner(db: Session, user_id: str, artifact: Artifact | None) -> Artifact:
    if artifact is None:
        raise HTTPException(status_code=404, detail="Artefacto no encontrado")
    project = db.get(Project, artifact.project_id)
    if project is None or project.owner_id != user_id:
        raise HTTPException(status_code=404, detail="Artefacto no encontrado")
    return artifact


def _artifact_read(db: Session, artifact: Artifact, include_content: bool) -> ArtifactRead:
    path = get_settings().data_dir / artifact.path
    content = None
    if include_content and artifact.format in TEXT_FORMATS and path.is_file():
        content = path.read_text(encoding="utf-8")
    if artifact.type == "slide_deck":
        renders = sorted(available_renders(path))
    else:
        renders = ["pdf"] if artifact.type == "lesson_content" else []
    versions = db.scalars(
        select(Artifact)
        .where(
            Artifact.project_id == artifact.project_id,
            Artifact.logical_key == artifact.logical_key,
        )
        .order_by(Artifact.version.desc())
    ).all()
    return ArtifactRead(
        id=artifact.id,
        project_id=artifact.project_id,
        type=artifact.type,
        format=artifact.format,
        title=artifact.title,
        logical_key=artifact.logical_key,
        version=artifact.version,
        is_selected=artifact.is_selected,
        metadata=artifact_metadata(artifact),
        created_by_job_id=artifact.created_by_job_id,
        created_at=artifact.created_at,
        content=content,
        renders=renders,
        versions=[
            ArtifactVersionRead(
                id=item.id,
                version=item.version,
                is_selected=item.is_selected,
                metadata=artifact_metadata(item),
                created_at=item.created_at,
            )
            for item in versions
        ],
    )


@router.get("/projects/{project_id}/artifacts", response_model=list[ArtifactRead])
def list_project_artifacts(project_id: str, user: CurrentUser, db: DB):
    project = db.get(Project, project_id)
    if project is None or project.owner_id != user.id:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    artifacts = db.scalars(
        select(Artifact)
        .where(Artifact.project_id == project_id, Artifact.is_selected.is_(True))
        .order_by(Artifact.created_at.desc(), Artifact.id.desc())
    ).all()
    return [_artifact_read(db, a, include_content=False) for a in artifacts]


@router.get("/artifacts/{artifact_id}", response_model=ArtifactRead)
def get_artifact(artifact_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    return _artifact_read(db, artifact, include_content=True)


@router.get("/artifacts/{artifact_id}/images/{image_id}")
def get_slide_image(artifact_id: str, image_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if artifact.type != "slide_deck":
        raise HTTPException(status_code=404, detail="Imagen no encontrada")
    image = next(
        (
            item
            for item in artifact_metadata(artifact).get("images", [])
            if item.get("id") == image_id and item.get("path")
        ),
        None,
    )
    path = _safe_stored_asset(str(image["path"])) if image else None
    if path is None or not path.is_file():
        raise HTTPException(status_code=404, detail="Imagen no encontrada")
    return FileResponse(
        path,
        media_type=image.get("media_type") or "application/octet-stream",
        filename=path.name,
        content_disposition_type="inline",
    )


@router.post(
    "/artifacts/{artifact_id}/images/{image_id}/regenerate",
    response_model=ArtifactRead,
)
def regenerate_slide_image(
    artifact_id: str,
    image_id: str,
    body: SlideImageRegenerate,
    user: CurrentUser,
    db: DB,
):
    """Regenerate one image and save the modified deck as a new immutable version."""
    original = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if original.type != "slide_deck":
        raise HTTPException(status_code=422, detail="El artefacto no es un deck de slides")
    metadata = artifact_metadata(original)
    original_image = next(
        (item for item in metadata.get("images", []) if item.get("id") == image_id),
        None,
    )
    if original_image is None:
        raise HTTPException(status_code=404, detail="Imagen no encontrada")

    model = original_image.get("model") or metadata.get("image_generation", {}).get(
        "model", DEFAULT_IMAGE_MODEL
    )
    style = original_image.get("style") or metadata.get("image_generation", {}).get(
        "style", DEFAULT_IMAGE_STYLE
    )
    style_prompt = original_image.get("style_prompt") or resolve_style_prompt(
        style,
        metadata.get("image_generation", {}).get("style_prompt", ""),
    )
    seed = original_image.get("seed")
    if seed is None:
        seed = consistency_seed(original.logical_key, style_prompt)
    resolved_prompt = (
        f"{body.prompt.strip()}\n\nSTYLE SYSTEM:\n{style_prompt}\n\n{CONSISTENCY_PROMPT}"
    )
    try:
        generated = generate_image(
            resolved_prompt,
            api_key=get_settings().openrouter_api_key,
            model=model,
            orientation=metadata.get("orientation", "horizontal"),
            seed=seed,
        )
    except (ImageGenerationError, ValueError) as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    settings = get_settings()
    original_path = settings.data_dir / original.path
    markdown = original_path.read_text(encoding="utf-8")
    markdown, cloned_metadata, output_dir, storage_asset_dir = _clone_slide_assets(
        markdown,
        metadata,
        original.project_id,
        suffix="regen",
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    cloned_image = next(
        item for item in cloned_metadata.get("images", []) if item.get("id") == image_id
    )
    previous_path = _safe_stored_asset(str(cloned_image.get("path", "")))
    extension = media_extension(generated.media_type, generated.content)
    filename = f"{image_id}{extension}"
    target = output_dir / filename
    target.write_bytes(generated.content)
    if previous_path is not None and previous_path != target:
        previous_path.unlink(missing_ok=True)
    markdown_path = f"{output_dir.name}/{filename}"
    requested_layout = cloned_image.get(
        "requested_layout", cloned_image.get("layout", "right")
    )
    _, effective_layout = normalize_automatic_image_layout(
        cloned_image.get("effective_layout", cloned_image.get("layout", "right"))
    )
    try:
        markdown = replace_generated_image(
            markdown,
            image_id,
            markdown_path,
            layout=effective_layout,
            alt=cloned_image.get("alt", "Ilustración generada"),
            orientation=cloned_metadata.get("orientation", "horizontal"),
        )
    except ValueError as exc:
        shutil.rmtree(output_dir, ignore_errors=True)
        raise HTTPException(status_code=409, detail=str(exc)) from exc

    cloned_image.update(
        {
            "prompt": body.prompt.strip(),
            "model": model,
            "style": style,
            "style_prompt": style_prompt,
            "seed": seed,
            "path": f"{storage_asset_dir}/{filename}",
            "markdown_path": markdown_path,
            "media_type": generated.media_type,
            "cost_usd": generated.cost_usd,
            "status": "generated",
            "requested_layout": requested_layout,
            "effective_layout": effective_layout,
            "layout": effective_layout,
        }
    )
    cloned_image.pop("error", None)
    generation = cloned_metadata.setdefault("image_generation", {})
    generation.update(
        {
            "enabled": True,
            "model": model,
            "style": style,
            "generated": sum(
                item.get("status") == "generated"
                for item in cloned_metadata.get("images", [])
            ),
            "attempted": len(cloned_metadata.get("images", [])),
            "generation_cost_usd": generated.cost_usd or 0.0,
            "regenerated_image_id": image_id,
        }
    )

    relative = (
        f"artifacts/{original.project_id}/slide_deck-image-edit-{uuid.uuid4().hex[:10]}.md"
    )
    path = settings.data_dir / relative
    try:
        cloned_metadata, _rendered = _prepare_slide_version(
            path,
            markdown,
            cloned_metadata,
        )
    except Exception:
        _remove_slide_assets(cloned_metadata, original.project_id)
        raise
    artifact = add_artifact_version(
        db,
        project_id=original.project_id,
        type_="slide_deck",
        format_="markdown",
        title=original.title,
        path=relative,
        metadata=cloned_metadata,
    )
    db.commit()
    db.refresh(artifact)
    usage_record_id = record_usage(
        project_id=artifact.project_id,
        agent="slides",
        operation="image",
        provider="openrouter",
        model=model,
        image_count=1,
        cost_usd=generated.cost_usd,
        artifact_id=artifact.id,
        work_unit_key=f"artifact-image:{artifact.id}:{image_id}",
        idempotency_key=f"artifact-image:{artifact.id}:{image_id}",
        metadata={"seed": seed, "style": style, "regeneration": True},
        emit_event=False,
    )
    if usage_record_id:
        attach_usage_to_artifact(
            db, artifact, usage_record_ids=[usage_record_id]
        )
        db.commit()
        db.refresh(artifact)
    logger.info(
        "Slide image regenerated as a new artifact version",
        extra={
            "artifact_id": artifact.id,
            "source_artifact_id": original.id,
            "project_id": original.project_id,
            "image_id": image_id,
            "artifact_version": artifact.version,
        },
    )
    return _artifact_read(db, artifact, include_content=True)


def _normalize_palette_request(palette: dict[str, str]) -> dict[str, str]:
    try:
        return normalize_palette(palette)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _cleanup_palette_files(paths: list[Path], metadata: list[dict], project_id: str) -> None:
    for path in paths:
        renders = available_renders(path)
        for candidate in [path, render_manifest_path(path), *renders.values()]:
            Path(candidate).unlink(missing_ok=True)
    for item in metadata:
        _remove_slide_assets(item, project_id)


@router.post("/artifacts/{artifact_id}/palette/preview")
def preview_slide_palette(
    artifact_id: str,
    body: SlidePalettePreview,
    user: CurrentUser,
    db: DB,
):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if artifact.type != "slide_deck":
        raise HTTPException(status_code=422, detail="El artefacto no es un deck de slides")
    if not marp_available():
        raise HTTPException(status_code=503, detail="Marp CLI no está disponible")
    palette = _normalize_palette_request(body.palette)
    source = get_settings().data_dir / artifact.path
    if not source.is_file():
        raise HTTPException(status_code=404, detail="Fichero no encontrado")
    preview = source.with_name(f".palette-preview-{uuid.uuid4().hex[:8]}.md")
    _preview_metadata, rendered = _prepare_slide_version(
        preview,
        apply_slide_palette(source.read_text(encoding="utf-8"), palette),
        artifact_metadata(artifact),
    )
    try:
        html = rendered.get("html")
        if not html:
            raise HTTPException(status_code=503, detail="No se pudo generar la preview")
        return Response(content=Path(html).read_bytes(), media_type="text/html")
    finally:
        _cleanup_palette_files([preview], [], artifact.project_id)


@router.post(
    "/artifacts/{artifact_id}/palette",
    response_model=list[ArtifactRead],
)
def apply_artifact_palette(
    artifact_id: str,
    body: SlidePaletteApply,
    user: CurrentUser,
    db: DB,
):
    """Create self-contained deck versions with new canonical palette renders."""
    original = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if original.type != "slide_deck":
        raise HTTPException(status_code=422, detail="El artefacto no es un deck de slides")
    if not marp_available():
        raise HTTPException(status_code=503, detail="Marp CLI no está disponible")
    palette = _normalize_palette_request(body.palette)
    if body.scope == "project":
        targets = _selected_artifacts(db, original.project_id, "slide_deck")
    else:
        targets = [original]
    if not targets:
        raise HTTPException(status_code=409, detail="No hay decks activos que actualizar")

    settings = get_settings()
    for target in targets:
        if not (settings.data_dir / target.path).is_file():
            raise HTTPException(
                status_code=409,
                detail=f"Falta el fichero del deck «{target.title}»",
            )

    prepared: list[tuple[Artifact, str, Path, dict]] = []
    created_paths: list[Path] = []
    cloned_metadata: list[dict] = []
    try:
        for target in targets:
            source = settings.data_dir / target.path
            markdown = source.read_text(encoding="utf-8")
            current_metadata = artifact_metadata(target)
            markdown, metadata, _asset_dir, _storage_dir = _clone_slide_assets(
                markdown,
                current_metadata,
                target.project_id,
                suffix="palette",
            )
            warnings = palette_warnings(palette)
            metadata.update(
                {
                    "source_artifact_id": target.id,
                    "palette_previous": normalize_palette(
                        current_metadata.get("slide_palette")
                        or DEFAULT_SLIDE_PALETTE
                    ),
                    "slide_palette": palette,
                    "palette_name": palette_preset_name(palette),
                    "palette_contrast": palette_contrast(palette),
                    "palette_warnings": warnings,
                    "version_reason": "palette_edit",
                }
            )
            relative = (
                f"artifacts/{target.project_id}/slide_deck-palette-"
                f"{uuid.uuid4().hex[:10]}.md"
            )
            output = settings.data_dir / relative
            output.parent.mkdir(parents=True, exist_ok=True)
            created_paths.append(output)
            cloned_metadata.append(metadata)
            metadata, _rendered = _prepare_slide_version(
                output,
                apply_slide_palette(markdown, palette),
                metadata,
            )
            cloned_metadata[-1] = metadata
            prepared.append((target, relative, output, metadata))

        created: list[Artifact] = []
        for target, relative, _output, metadata in prepared:
            created.append(
                add_artifact_version(
                    db,
                    project_id=target.project_id,
                    type_="slide_deck",
                    format_="markdown",
                    title=target.title,
                    path=relative,
                    metadata=metadata,
                )
            )
        db.commit()
        for artifact in created:
            db.refresh(artifact)
        logger.info(
            "Slide palette applied as new artifact versions",
            extra={
                "project_id": original.project_id,
                "source_artifact_id": original.id,
                "artifact_count": len(created),
                "palette_name": palette_preset_name(palette),
                "scope": body.scope,
            },
        )
        return [_artifact_read(db, artifact, include_content=True) for artifact in created]
    except Exception:
        db.rollback()
        _cleanup_palette_files(created_paths, cloned_metadata, original.project_id)
        raise


@router.patch("/artifacts/{artifact_id}", response_model=ArtifactRead)
def edit_artifact(
    artifact_id: str,
    body: ArtifactEdit,
    user: CurrentUser,
    db: DB,
):
    """Save edited text as the next immutable artifact version."""
    original = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if original.type == "audio":
        raise HTTPException(
            status_code=422,
            detail="Los manifiestos de audio son inmutables; regenera el agente de Audio",
        )
    if original.format not in TEXT_FORMATS:
        raise HTTPException(status_code=422, detail="Este formato no se puede editar como texto")
    if original.format == "json":
        try:
            json.loads(body.content)
        except json.JSONDecodeError as exc:
            raise HTTPException(
                status_code=422,
                detail=f"JSON invalido en la linea {exc.lineno}: {exc.msg}",
            ) from exc

    settings = get_settings()
    extension = (
        Path(original.path).suffix
        or {
            "markdown": ".md",
            "json": ".json",
            "text": ".txt",
        }[original.format]
    )
    relative = (
        f"artifacts/{original.project_id}/edit-{original.type}-{uuid.uuid4().hex[:10]}{extension}"
    )
    path = settings.data_dir / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    content = body.content
    metadata = artifact_metadata(original)
    cloned_slide_assets = False
    if original.type == "slide_deck" and (
        metadata.get("images") or isinstance(metadata.get("logo"), dict)
    ):
        content, metadata, _asset_dir, _storage_dir = _clone_slide_assets(
            content,
            metadata,
            original.project_id,
            suffix="edit",
        )
        cloned_slide_assets = True
    if original.type == "publication_package" and isinstance(
        metadata.get("thumbnail"), dict
    ):
        thumbnail = dict(metadata["thumbnail"])
        source_thumbnail = settings.data_dir / str(thumbnail.get("path") or "")
        if source_thumbnail.is_file():
            thumbnail_relative = (
                f"artifacts/{original.project_id}/edit-publication-"
                f"{uuid.uuid4().hex[:10]}-thumbnail{source_thumbnail.suffix or '.png'}"
            )
            thumbnail_target = settings.data_dir / thumbnail_relative
            thumbnail_target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source_thumbnail, thumbnail_target)
            thumbnail["path"] = thumbnail_relative
            metadata["thumbnail"] = thumbnail
    if original.type == "slide_deck":
        try:
            metadata, _rendered = _prepare_slide_version(path, content, metadata)
        except Exception:
            if cloned_slide_assets:
                _remove_slide_assets(metadata, original.project_id)
            raise
    else:
        path.write_text(content, encoding="utf-8")
    edited = add_artifact_version(
        db,
        project_id=original.project_id,
        type_=original.type,
        format_=original.format,
        title=original.title,
        path=relative,
        metadata=metadata,
    )
    db.commit()
    db.refresh(edited)
    logger.info(
        "Artifact version created from web editor",
        extra={
            "artifact_id": edited.id,
            "project_id": edited.project_id,
            "artifact_type": edited.type,
            "artifact_version": edited.version,
        },
    )
    return _artifact_read(db, edited, include_content=True)


@router.post("/artifacts/{artifact_id}/select", response_model=ArtifactRead)
def select_artifact(artifact_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    select_artifact_version(db, artifact)
    db.commit()
    db.refresh(artifact)
    return _artifact_read(db, artifact, include_content=False)


@router.delete("/artifacts/{artifact_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_artifact(artifact_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    path = get_settings().data_dir / artifact.path
    metadata = artifact_metadata(artifact)
    renders = available_renders(path) if artifact.type == "slide_deck" else {}
    if artifact.is_selected:
        select_latest_remaining(db, artifact)
    db.delete(artifact)
    db.commit()
    for candidate in [path, render_manifest_path(path), *renders.values()]:
        try:
            Path(candidate).unlink(missing_ok=True)
        except OSError:
            # The DB deletion is authoritative; orphan cleanup is best effort.
            pass
    if artifact.type == "slide_deck":
        _remove_slide_assets(metadata, artifact.project_id)
    if artifact.type == "publication_package" and isinstance(
        metadata.get("thumbnail"), dict
    ):
        thumbnail_relative = str(metadata["thumbnail"].get("path") or "")
        if thumbnail_relative:
            thumbnail_path = get_settings().data_dir / thumbnail_relative
            thumbnail_path.unlink(missing_ok=True)
    if artifact.type == "audio":
        data_root = get_settings().data_dir.resolve()
        for segment in metadata.get("segments") or []:
            if not isinstance(segment, dict) or not segment.get("path"):
                continue
            candidate = (data_root / str(segment["path"])).resolve()
            if candidate.is_relative_to(data_root):
                candidate.unlink(missing_ok=True)


@router.get("/artifacts/{artifact_id}/audio/{segment_index}")
def get_audio_segment(
    artifact_id: str,
    segment_index: int,
    user: CurrentUser,
    db: DB,
):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if artifact.type != "audio":
        raise HTTPException(status_code=404, detail="Segmento de audio no encontrado")
    segments = artifact_metadata(artifact).get("segments") or []
    segment = next(
        (
            item
            for item in segments
            if isinstance(item, dict) and item.get("index") == segment_index
        ),
        None,
    )
    if segment is None or not segment.get("path"):
        raise HTTPException(status_code=404, detail="Segmento de audio no encontrado")
    data_root = get_settings().data_dir.resolve()
    path = (data_root / str(segment["path"])).resolve()
    if not path.is_relative_to(data_root) or not path.is_file():
        raise HTTPException(status_code=404, detail="Segmento de audio no encontrado")
    expected_sha256 = str(segment.get("sha256") or "")
    if expected_sha256 and hashlib.sha256(path.read_bytes()).hexdigest() != expected_sha256:
        raise HTTPException(status_code=409, detail="El segmento de audio está dañado")
    return FileResponse(
        path,
        media_type=str(segment.get("mime_type") or "application/octet-stream"),
        filename=f"{artifact.title}-segmento-{segment_index}{path.suffix}",
    )


@router.get("/artifacts/{artifact_id}/thumbnail")
def get_publication_thumbnail(artifact_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    if artifact.type != "publication_package":
        raise HTTPException(status_code=404, detail="Miniatura no encontrada")
    thumbnail = artifact_metadata(artifact).get("thumbnail") or {}
    relative_path = thumbnail.get("path") if isinstance(thumbnail, dict) else None
    path = get_settings().data_dir / relative_path if relative_path else None
    if path is None or not path.is_file():
        raise HTTPException(status_code=404, detail="Miniatura no encontrada")
    return FileResponse(
        path,
        media_type=str(thumbnail.get("media_type") or "image/png"),
        filename=f"{artifact.title}-miniatura{path.suffix or '.png'}",
    )


@router.get("/artifacts/{artifact_id}/download")
def download_artifact(artifact_id: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    path = get_settings().data_dir / artifact.path
    if not path.is_file():
        raise HTTPException(status_code=404, detail="Fichero no encontrado")
    if artifact.type == "slide_deck":
        return Response(
            content=normalize_marp_canvas(path.read_text(encoding="utf-8")),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{path.name}"'},
        )
    return FileResponse(path, filename=path.name)


@router.get("/artifacts/{artifact_id}/render/{fmt}")
def get_artifact_render(artifact_id: str, fmt: str, user: CurrentUser, db: DB):
    artifact = _check_owner(db, user.id, db.get(Artifact, artifact_id))
    path = get_settings().data_dir / artifact.path
    if artifact.type == "lesson_content":
        if fmt != "pdf" or not path.is_file():
            raise HTTPException(status_code=404, detail=f"Render {fmt} no disponible")
        content = lessons_pdf(
            [
                LessonDocument(
                    title=artifact.title or "Leccion",
                    markdown=path.read_text(encoding="utf-8"),
                )
            ]
        )
        logger.info(
            "Lesson PDF exported",
            extra={"artifact_id": artifact.id, "project_id": artifact.project_id},
        )
        return _download_response(
            content,
            "application/pdf",
            _safe_filename(artifact.title, "leccion") + ".pdf",
        )

    if artifact.type != "slide_deck" or fmt not in RENDER_MEDIA_TYPES:
        raise HTTPException(status_code=404, detail="Formato desconocido")
    renders = _ensure_slide_renders(path)
    if fmt not in renders:
        raise HTTPException(status_code=404, detail=f"Render {fmt} no disponible")
    logger.info(
        "Slide render downloaded",
        extra={
            "artifact_id": artifact.id,
            "project_id": artifact.project_id,
            "render_format": fmt,
        },
    )
    return FileResponse(
        renders[fmt],
        media_type=RENDER_MEDIA_TYPES[fmt],
        filename=f"{artifact.title}.{fmt}",
        content_disposition_type="inline" if fmt == "html" else "attachment",
    )


@router.get("/projects/{project_id}/exports/lessons.pdf")
def export_lessons_pdf(project_id: str, user: CurrentUser, db: DB):
    project = _check_project(db, user.id, project_id)
    documents = _artifact_documents(_selected_artifacts(db, project_id, "lesson_content"))
    if not documents:
        raise HTTPException(status_code=404, detail="No hay lecciones activas para exportar")
    content = lessons_pdf(documents)
    logger.info(
        "Combined lessons PDF exported",
        extra={"project_id": project_id, "artifact_count": len(documents)},
    )
    return _download_response(
        content,
        "application/pdf",
        _safe_filename(project.title, "curso") + "-lecciones.pdf",
    )


@router.get("/projects/{project_id}/exports/lessons.zip")
def export_lessons_zip(project_id: str, user: CurrentUser, db: DB):
    project = _check_project(db, user.id, project_id)
    documents = _artifact_documents(_selected_artifacts(db, project_id, "lesson_content"))
    if not documents:
        raise HTTPException(status_code=404, detail="No hay lecciones activas para exportar")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for index, document in enumerate(documents, start=1):
            name = f"{index:02d}-{_safe_filename(document.title, 'leccion')}.pdf"
            archive.writestr(name, lessons_pdf([document]))
    logger.info(
        "Lesson PDF ZIP exported",
        extra={"project_id": project_id, "artifact_count": len(documents)},
    )
    return _download_response(
        buffer.getvalue(),
        "application/zip",
        _safe_filename(project.title, "curso") + "-lecciones.zip",
    )


@router.get("/projects/{project_id}/exports/slides.pdf")
def export_slides_pdf(project_id: str, user: CurrentUser, db: DB):
    project = _check_project(db, user.id, project_id)
    artifacts = _selected_artifacts(db, project_id, "slide_deck")
    if not artifacts:
        raise HTTPException(status_code=404, detail="No hay slides activas para exportar")
    writer = PdfWriter()
    settings = get_settings()
    for artifact in artifacts:
        renders = _ensure_slide_renders(settings.data_dir / artifact.path)
        pdf_path = renders.get("pdf")
        if pdf_path is None:
            raise HTTPException(
                status_code=503,
                detail="No se pudieron renderizar todas las slides como PDF",
            )
        writer.append(PdfReader(pdf_path))
    buffer = io.BytesIO()
    writer.write(buffer)
    logger.info(
        "Combined slides PDF exported",
        extra={"project_id": project_id, "artifact_count": len(artifacts)},
    )
    return _download_response(
        buffer.getvalue(),
        "application/pdf",
        _safe_filename(project.title, "curso") + "-slides.pdf",
    )


@router.get("/projects/{project_id}/exports/slides.pptx")
def export_slides_pptx(project_id: str, user: CurrentUser, db: DB):
    project = _check_project(db, user.id, project_id)
    artifacts = _selected_artifacts(db, project_id, "slide_deck")
    if not artifacts:
        raise HTTPException(status_code=404, detail="No hay slides activas para exportar")
    orientations = {
        artifact_metadata(artifact).get("orientation", "horizontal")
        for artifact in artifacts
    }
    if len(orientations) > 1:
        raise HTTPException(
            status_code=422,
            detail=(
                "No se puede crear un único PPTX con slides horizontales y verticales. "
                "Selecciona decks de una sola orientación o descarga el PDF o ZIP."
            ),
        )
    content = _render_combined_slides(artifacts, "pptx")
    logger.info(
        "Combined slides PPTX exported",
        extra={"project_id": project_id, "artifact_count": len(artifacts)},
    )
    return _download_response(
        content,
        RENDER_MEDIA_TYPES["pptx"],
        _safe_filename(project.title, "curso") + "-slides.pptx",
    )


@router.get("/projects/{project_id}/exports/slides.zip")
def export_slides_zip(project_id: str, user: CurrentUser, db: DB):
    project = _check_project(db, user.id, project_id)
    artifacts = _selected_artifacts(db, project_id, "slide_deck")
    if not artifacts:
        raise HTTPException(status_code=404, detail="No hay slides activas para exportar")
    settings = get_settings()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        if any(
            artifact_metadata(artifact).get("orientation") == "vertical"
            for artifact in artifacts
        ):
            archive.write(vertical_theme_path(), "factory-vertical.css")
        for index, artifact in enumerate(artifacts, start=1):
            source = settings.data_dir / artifact.path
            renders = _ensure_slide_renders(source)
            stem = f"{index:02d}-{_safe_filename(artifact.title, 'slides')}"
            markdown = normalize_marp_canvas(source.read_text(encoding="utf-8"))
            metadata = artifact_metadata(artifact)
            portable_assets = list(metadata.get("images", []))
            if isinstance(metadata.get("logo"), dict):
                portable_assets.append(metadata["logo"])
            for image in portable_assets:
                stored_path = image.get("path")
                markdown_path = image.get("markdown_path")
                image_path = _safe_stored_asset(str(stored_path or ""))
                if image_path is None or not image_path.is_file() or not markdown_path:
                    continue
                archive_path = f"{stem}-assets/{image_path.name}"
                markdown = markdown.replace(str(markdown_path), archive_path)
                archive.write(image_path, archive_path)
            archive.writestr(f"{stem}.md", markdown)
            for fmt in ("pdf", "pptx"):
                if fmt in renders:
                    archive.write(renders[fmt], f"{stem}.{fmt}")
    logger.info(
        "Slides ZIP exported",
        extra={"project_id": project_id, "artifact_count": len(artifacts)},
    )
    return _download_response(
        buffer.getvalue(),
        "application/zip",
        _safe_filename(project.title, "curso") + "-slides.zip",
    )


@router.post(
    "/projects/{project_id}/artifacts",
    response_model=ArtifactRead,
    status_code=status.HTTP_201_CREATED,
)
async def upload_artifact(
    project_id: str,
    user: CurrentUser,
    db: DB,
    file: Annotated[UploadFile, File()],
    type: Annotated[str, Form()],
    title: Annotated[str, Form()] = "",
):
    """Entry point for external material (e.g. your own slides or notes)."""
    project = db.get(Project, project_id)
    if project is None or project.owner_id != user.id:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    if type not in UPLOADABLE_TYPES:
        raise HTTPException(
            status_code=422,
            detail=f"Tipo no subible. Permitidos: {', '.join(sorted(UPLOADABLE_TYPES))}",
        )
    settings = get_settings()
    format_ = UPLOADABLE_TYPES[type]
    ext = {"markdown": "md", "json": "json"}[format_]
    rel_path = f"artifacts/{project_id}/upload-{type}-{uuid.uuid4().hex[:8]}.{ext}"
    abs_path = settings.data_dir / rel_path
    abs_path.parent.mkdir(parents=True, exist_ok=True)
    data = await file.read()
    if type == "slide_deck" and (file.filename or "").lower().endswith(".pptx"):
        try:
            data = pptx_to_marp(data).encode("utf-8")
        except Exception as exc:
            raise HTTPException(
                status_code=422, detail=f"No se pudo convertir el PPTX: {exc}"
            ) from exc
    if format_ == "json":
        try:
            json.loads(data.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise HTTPException(status_code=422, detail="El fichero JSON no es valido") from exc
    metadata = (
        {
            "orientation": "horizontal",
            "width": 1920,
            "height": 1080,
        }
        if type == "slide_deck"
        else None
    )
    if type == "slide_deck":
        try:
            markdown = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=422, detail="Las slides deben usar texto UTF-8 válido"
            ) from exc
        metadata, _rendered = _prepare_slide_version(abs_path, markdown, metadata or {})
    else:
        abs_path.write_bytes(data)
    artifact = add_artifact_version(
        db,
        project_id=project_id,
        type_=type,
        format_=format_,
        title=title or (file.filename or type),
        path=rel_path,
        metadata=metadata,
    )
    db.commit()
    db.refresh(artifact)
    logger.info(
        "Artifact uploaded",
        extra={
            "artifact_id": artifact.id,
            "project_id": project_id,
            "artifact_type": type,
            "artifact_version": artifact.version,
        },
    )
    return _artifact_read(db, artifact, include_content=False)
